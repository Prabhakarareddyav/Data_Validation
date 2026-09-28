"""mRNA Data Validator - compares an Excel source against a CSV target (e.g. Snowflake export)."""
import io, os
from datetime import datetime
import pandas as pd
from flask import Flask, request, jsonify, send_file, send_from_directory
from flask_cors import CORS
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

app = Flask(__name__, static_folder="static")
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024  # 200 MB uploads

ALWAYS_IGNORE = {"RECORD_COUNT"}


def log(msg):
    return f"[{datetime.now():%H:%M:%S}] {msg}"


def norm(v):
    """Normalise a value so 1, 1.0, ' 1 ' and NaN compare sensibly."""
    if pd.isna(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    v = str(v).strip()
    try:
        f = float(v)
        return str(int(f)) if f.is_integer() else str(f)
    except ValueError:
        return v


def read_any(file, sheet, header):
    name = (file.filename or "").lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        return pd.read_excel(file, sheet_name=sheet, header=header)
    return pd.read_csv(file, header=header)


@app.get("/")
def index():
    return send_from_directory("static", "index.html")


@app.get("/api/health")
def health():
    return jsonify(status="ok", version="2.0.0", timestamp=datetime.now().isoformat())


@app.post("/api/validate")
def validate():
    logs = []
    try:
        src, tgt = request.files.get("source_file"), request.files.get("target_file")
        if not src or not tgt:
            return jsonify(status="error", message="Upload both a source and a target file."), 400

        f = request.form
        sheet = f.get("source_sheet") or 0
        header = int(f.get("header_row", 3))
        key = f.get("key_column", "Antigen_ID").strip().upper()
        skip_first = f.get("skip_first_column", "true") == "true"
        ignore = ALWAYS_IGNORE | {c.strip().upper() for c in f.get("ignore_columns", "").split(",") if c.strip()}

        logs.append(log(f"Source: {src.filename} | Target: {tgt.filename}"))
        s = read_any(src, sheet, header)
        if skip_first:
            s = s.iloc[:, 1:]
        t = read_any(tgt, 0, 0)
        s.columns = s.columns.astype(str).str.strip().str.upper()
        t.columns = t.columns.astype(str).str.strip().str.upper()

        for name, df in (("source", s), ("target", t)):
            if key not in df.columns:
                return jsonify(status="error", log=logs,
                    message=f"Key column '{key}' not found in {name}. Columns found: {', '.join(df.columns[:15])}"), 400

        cols = [c for c in s.columns if c not in ignore and not c.startswith("UNNAMED")]
        missing_cols = [c for c in cols if c not in t.columns]
        if missing_cols:
            logs.append(log(f"Columns absent in target (reported as mismatches): {', '.join(missing_cols)}"))
        logs.append(log(f"Columns compared: {len(cols)} | Source rows: {len(s)} | Target rows: {len(t)}"))

        s[key] = s[key].map(norm); t[key] = t[key].map(norm)
        s, t = s[s[key] != ""], t[t[key] != ""]
        dup_s, dup_t = int(s[key].duplicated().sum()), int(t[key].duplicated().sum())
        if dup_s or dup_t:
            logs.append(log(f"Duplicate keys (first kept) - source: {dup_s}, target: {dup_t}"))
        s = s.drop_duplicates(key).set_index(key)
        t = t.drop_duplicates(key).set_index(key)

        sk, tk = set(s.index), set(t.index)
        common = sorted(sk & tk)
        missing, extra = sorted(sk - tk), sorted(tk - sk)
        logs.append(log(f"Matching: {len(common)} | Missing in target: {len(missing)} | Extra in target: {len(extra)}"))

        mismatches = []
        for col in (c for c in cols if c != key):
            a = s.loc[common, col].map(norm)
            b = t.loc[common, col].map(norm) if col in t.columns else pd.Series("<column missing>", index=common)
            for k in a.index[a.values != b.values]:
                mismatches.append({"Key": k, "Column": col, "Source_Value": a[k], "Target_Value": b[k]})
        logs.append(log(f"Mismatches: {len(mismatches)}" if mismatches else "All matching records are identical."))

        return jsonify(status="ok", log=logs, mismatches=mismatches, summary={
            "source_rows": len(s), "target_rows": len(t), "matching_keys": len(common),
            "missing_in_target": len(missing), "extra_in_target": len(extra),
            "total_mismatches": len(mismatches), "columns_compared": len(cols),
            "missing_keys_list": missing, "extra_keys_list": extra})
    except Exception as e:
        logs.append(log(f"ERROR: {e}"))
        return jsonify(status="error", message=str(e), log=logs), 500


@app.post("/api/export")
def export():
    d = request.get_json()
    sm, mm = d.get("summary", {}), d.get("mismatches", [])
    wb = Workbook()
    head = PatternFill("solid", fgColor="4B1F6F")
    bad = PatternFill("solid", fgColor="FBE3E0")

    def sheet(ws, headers, rows, shade=None):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font, c.alignment = head, Font(bold=True, color="FFFFFF"), Alignment(horizontal="center")
        for r in rows:
            ws.append(r)
            if shade:
                for c in ws[ws.max_row]:
                    c.fill = shade
        for col in ws.columns:
            w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(w + 3, 50)

    labels = [("Validation date", datetime.now().strftime("%Y-%m-%d %H:%M")),
              ("Source rows", sm.get("source_rows")), ("Target rows", sm.get("target_rows")),
              ("Columns compared", sm.get("columns_compared")), ("Matching keys", sm.get("matching_keys")),
              ("Missing in target", sm.get("missing_in_target")), ("Extra in target", sm.get("extra_in_target")),
              ("Total mismatches", sm.get("total_mismatches"))]
    sheet(wb.active, ["Metric", "Value"], labels); wb.active.title = "Summary"
    sheet(wb.create_sheet("Mismatches"), ["Key", "Column", "Source_Value", "Target_Value"],
          [[m["Key"], m["Column"], m["Source_Value"], m["Target_Value"]] for m in mm], bad)
    sheet(wb.create_sheet("Missing Keys"), ["Keys missing in target"], [[k] for k in sm.get("missing_keys_list", [])])
    sheet(wb.create_sheet("Extra Keys"), ["Extra keys in target"], [[k] for k in sm.get("extra_keys_list", [])])

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="validation_report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    print(f"mRNA Data Validator running at http://localhost:{port}")
    app.run(host="0.0.0.0", port=port)
